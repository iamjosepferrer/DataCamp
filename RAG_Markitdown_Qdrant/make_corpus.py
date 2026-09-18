"""Generate the coffee-shop knowledge base (4 mixed-format files)."""
from pathlib import Path
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

KB = Path("knowledge_base")
KB.mkdir(exist_ok=True)

# 1. catalog.pdf: a pricing table
styles = getSampleStyleSheet()
doc = SimpleDocTemplate(str(KB / "catalog.pdf"), pagesize=A4)
rows = [
    ["Product", "8oz", "12oz", "16oz", "Origin"],
    ["House Blend", "$3.20", "$3.80", "$4.40", "Brazil / Colombia"],
    ["Single Origin Ethiopia", "$3.90", "$4.60", "$5.20", "Yirgacheffe"],
    ["Cold Brew", "n/a", "$4.20", "$4.90", "House Blend base"],
    ["Decaf Swiss Water", "$3.20", "$3.80", "$4.40", "Colombia"],
    ["Oat Milk Latte", "$4.50", "$5.10", "$5.70", "House Blend base"],
]
t = Table(rows, colWidths=[150, 55, 55, 55, 130])
t.setStyle(TableStyle([
    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
    ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
]))
doc.build([
    Paragraph("Bean & Bloom Coffee: Product Catalog 2026", styles["Title"]),
    Paragraph("Drinks", styles["Heading2"]),
    Paragraph("All prices include VAT. Sizes are in fluid ounces.", styles["Normal"]),
    Spacer(1, 12), t, Spacer(1, 12),
    Paragraph("Whole Bean Bags", styles["Heading2"]),
    Paragraph("250g bags of any roast cost $12.50. 1kg bags cost $42.00. "
              "Subscriptions get 10% off every bag.", styles["Normal"]),
])

# 2. returns_policy.docx: headed sections
d = Document()
d.add_heading("Bean & Bloom Returns Policy", 0)
d.add_heading("Drinks", 1)
d.add_paragraph("Remade free of charge if you tell us before leaving the counter. "
                "No refunds on drinks after they leave the shop.")
d.add_heading("Whole Bean Bags", 1)
d.add_paragraph("Unopened bags can be returned within 14 days with a receipt for a full refund. "
                "Opened bags can be exchanged, not refunded.")
d.add_heading("Damaged Bags", 1)
d.add_paragraph("If a bag arrives torn or the valve has failed, email hello@beanandbloom.example "
                "with a photo within 48 hours. We ship a replacement at no cost.")
d.add_heading("Equipment", 1)
d.add_paragraph("Grinders and brewers carry a 12-month warranty. Returns require original packaging.")
d.save(KB / "returns_policy.docx")

# 3. q2_sales.xlsx: one sheet
wb = Workbook()
ws = wb.active
ws.title = "Q2 2026"
ws.append(["Month", "Drinks revenue", "Bean bag revenue", "Equipment revenue", "Total"])
for m, a, b, c in [("April", 41200, 12800, 3100), ("May", 44650, 13900, 2750), ("June", 47300, 15200, 4600)]:
    ws.append([m, a, b, c, a + b + c])
ws.append(["Q2 total", 133150, 41900, 10450, 185500])
wb.save(KB / "q2_sales.xlsx")

# 4. kickoff.pptx: slides + speaker notes
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Q3 Kickoff: Bean & Bloom"
s.placeholders[1].text = "Goals for the quarter\nLaunch cold brew cans\nOpen second location in Rotterdam"
s.notes_slide.notes_text_frame.text = "Cold brew cans target launch is 15 October. Budget approved at $18k."
s2 = prs.slides.add_slide(prs.slide_layouts[1])
s2.shapes.title.text = "Risks"
s2.placeholders[1].text = "Green coffee prices up 22% year over year\nStaffing for the second shop"
s2.notes_slide.notes_text_frame.text = "If bean prices keep climbing we revisit the catalog in November."
prs.save(KB / "kickoff.pptx")

for p in sorted(KB.iterdir()):
    print(f"{p.name:22s} {p.stat().st_size:>7,} bytes")
